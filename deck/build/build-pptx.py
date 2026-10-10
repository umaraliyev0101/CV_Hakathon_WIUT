#!/usr/bin/env python3
"""Clio Teardown — JSON slide specs -> native PowerPoint.

Text layouts become real, editable PowerPoint text.
Chart and diagram slides are rendered by Chrome and placed as full-bleed images,
so they look identical to the PDF.

Usage: build-pptx.py <slidesDir> <out.pptx> <shotsDir>
"""
import json
import os
import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

SLIDES_DIR, OUT, SHOTS = sys.argv[1], sys.argv[2], sys.argv[3]

# ---- design tokens, mirrored from design.css ----
PAPER = RGBColor(0xFB, 0xFA, 0xF8)
PAPER_ALT = RGBColor(0xF3, 0xF1, 0xEC)
INK = RGBColor(0x1A, 0x19, 0x17)
INK_SOFT = RGBColor(0x4A, 0x47, 0x41)
MUTED = RGBColor(0x7C, 0x78, 0x71)
HAIRLINE = RGBColor(0xDE, 0xDB, 0xD4)
CLAY = RGBColor(0xC2, 0x49, 0x2D)
CLAY_LT = RGBColor(0xE8, 0x83, 0x6A)
POS = RGBColor(0x4A, 0x6B, 0x3D)
NEG = RGBColor(0xA3, 0x31, 0x1B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Helvetica Neue"

ACCENTS = ["#C2492D","#BC7A22","#6F7F32","#3F7D4E","#2E7D74","#2F6E93",
           "#4A56A0","#6E4C9B","#8E3F77","#B23A63","#55606E"]

def _rgb(h):
    h = h.lstrip("#"); return RGBColor(int(h[0:2],16), int(h[2:4],16), int(h[4:6],16))

def _mix(h, pct, other=(255,255,255)):
    """blend hex toward `other` by pct (0..1)"""
    h = h.lstrip("#"); r,g,b = int(h[0:2],16), int(h[2:4],16), int(h[4:6],16)
    f = lambda a,bb: int(round(a*(1-pct) + bb*pct))
    return RGBColor(f(r,other[0]), f(g,other[1]), f(b,other[2]))

class _A:
    """current part's hue, set per slide"""
    hex = ACCENTS[0]
    @property
    def base(self):  return _rgb(self.hex)
    @property
    def dk(self):    return _mix(self.hex, .22, (0,0,0))
    @property
    def t05(self):   return _mix(self.hex, .94)
    @property
    def t10(self):   return _mix(self.hex, .89)
    @property
    def t18(self):   return _mix(self.hex, .80)
ACC = _A()

PX = 13.3333 / 1280.0          # px -> inches
def I(px): return Inches(px * PX)
def P(px): return Pt(px * 0.75)  # css px -> points

PAD_L, PAD_T, PAD_R = 88, 72, 88
CONTENT_W = 1280 - PAD_L - PAD_R

prs = Presentation()
prs.slide_width = Inches(13.3333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

# ---------------- helpers ----------------

TOKEN = re.compile(r"\*\*(.+?)\*\*|`([^`]+?)`")

def put_runs(para, text, size, color, bold=False, font=FONT):
    """Write text into a paragraph, honouring **bold** and `code`."""
    pos = 0
    for m in TOKEN.finditer(text or ""):
        if m.start() > pos:
            r = para.add_run(); r.text = text[pos:m.start()]
            _style(r, size, color, bold, font)
        if m.group(1) is not None:
            r = para.add_run(); r.text = m.group(1)
            _style(r, size, INK if color in (INK_SOFT, MUTED) else color, True, font)
        else:
            r = para.add_run(); r.text = m.group(2)
            _style(r, size, color, bold, "Menlo")
        pos = m.end()
    if pos < len(text or ""):
        r = para.add_run(); r.text = text[pos:]
        _style(r, size, color, bold, font)
    if not (text or ""):
        r = para.add_run(); r.text = ""
        _style(r, size, color, bold, font)

def _style(run, size, color, bold, font):
    run.font.size = size; run.font.color.rgb = color
    run.font.bold = bold; run.font.name = font

def textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf

def line(slide, text, size, color, bold=False, space_after=0, space_before=0,
         tf=None, x=None, y=None, w=None, h=None, align=None, spacing=None):
    if tf is None:
        tf = textbox(slide, x, y, w, h)
        para = tf.paragraphs[0]
    else:
        para = tf.paragraphs[0] if (len(tf.paragraphs) == 1 and not tf.paragraphs[0].runs) else tf.add_paragraph()
    para.space_after = Pt(space_after); para.space_before = Pt(space_before)
    if align: para.alignment = align
    if spacing: para.line_spacing = spacing
    put_runs(para, text, size, color, bold)
    return tf

def bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color

def rule(slide, x, y, w, color=HAIRLINE, weight=1.0):
    from pptx.enum.shapes import MSO_SHAPE
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(x), I(y), I(w), Emu(int(weight * 9525)))
    s.fill.solid(); s.fill.fore_color.rgb = color; s.line.fill.background()
    s.shadow.inherit = False
    return s

def vrule(slide, x, y, h, color=HAIRLINE):
    from pptx.enum.shapes import MSO_SHAPE
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(x), I(y), Emu(9525), I(h))
    s.fill.solid(); s.fill.fore_color.rgb = color; s.line.fill.background()
    s.shadow.inherit = False
    return s

def strip(t):
    return TOKEN.sub(lambda m: m.group(1) or m.group(2), t or "")

def cell_text(c):
    return c.get("t", "") if isinstance(c, dict) else c

def furniture(slide, s, n, dark=False):
    c = RGBColor(0x8A, 0x87, 0x82) if dark else MUTED
    if s.get("source"):
        line(slide, s["source"], P(12.5), c, x=PAD_L, y=676, w=980, h=30)
    tf = textbox(slide, 1090, 676, 102, 30)
    para = tf.paragraphs[0]; para.alignment = PP_ALIGN.RIGHT
    put_runs(para, str(n), P(12.5), c)

def header(slide, s, dark=False):
    """eyebrow + title + subtitle. Returns y of the content area."""
    y = PAD_T
    sub_c = RGBColor(0xC8, 0xC5, 0xC0) if dark else INK_SOFT
    if s.get("eyebrow"):
        line(slide, s["eyebrow"].upper(), P(13), MUTED if not dark else RGBColor(0x9A, 0x97, 0x92),
             bold=True, x=PAD_L, y=y, w=CONTENT_W, h=24)
        y += 32
    if s.get("title"):
        t = s["title"]; size = 33 if len(strip(t)) > 52 else 40
        h = 2 if len(strip(t)) > 34 else 1
        line(slide, t, P(size), WHITE if dark else INK, bold=True,
             x=PAD_L, y=y - 6, w=min(CONTENT_W, 900), h=size * 1.3 * h, spacing=1.14)
        y += size * 1.2 * h + 4
    if s.get("subtitle"):
        line(slide, s["subtitle"], P(21), sub_c,
             x=PAD_L, y=y, w=min(CONTENT_W, 880), h=64, spacing=1.42)
        y += 34 + (30 if len(strip(s["subtitle"])) > 78 else 0)
    return y + 26

# ---------------- layouts ----------------

def L_title(slide, s):
    dark = s.get("theme") == "dark"
    bg(slide, INK if dark else PAPER)
    y = 168
    if s.get("eyebrow"):
        line(slide, s["eyebrow"].upper(), P(13), RGBColor(0x9A, 0x97, 0x92), bold=True,
             x=PAD_L, y=y, w=900, h=24); y += 36
    line(slide, s["title"], P(76), WHITE if dark else INK, bold=True,
         x=PAD_L, y=y, w=1000, h=180, spacing=1.04); y += 96
    if s.get("subtitle"):
        line(slide, s["subtitle"], P(25), RGBColor(0xC8, 0xC5, 0xC0) if dark else INK_SOFT,
             x=PAD_L, y=y, w=880, h=90, spacing=1.4); y += 92
    if s.get("meta"):
        tf = textbox(slide, PAD_L, y + 20, 900, 90)
        for i, m in enumerate(s["meta"]):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.space_after = Pt(5)
            put_runs(para, m, P(15), RGBColor(0x8A, 0x87, 0x82) if dark else MUTED)

def L_section(slide, s):
    bg(slide, ACC.base)
    y = 400
    if s.get("num"):
        line(slide, s["num"].upper(), P(15), RGBColor(0x9A, 0x97, 0x92), bold=True,
             x=PAD_L, y=y, w=600, h=26); y += 38
    line(slide, s["title"], P(64), WHITE, bold=True, x=PAD_L, y=y, w=1050, h=88, spacing=1.06)
    if s.get("subtitle"):
        line(slide, s["subtitle"], P(20), RGBColor(0xC8, 0xC5, 0xC0),
             x=PAD_L, y=y + 82, w=900, h=60, spacing=1.4)

def L_statement(slide, s):
    dark = s.get("theme") == "dark"
    bg(slide, INK if dark else ACC.t05)
    y = 190
    if s.get("eyebrow"):
        line(slide, s["eyebrow"].upper(), P(13), MUTED, bold=True, x=PAD_L, y=y, w=900, h=24)
        y += 34
    if s.get("big"):
        n = len(strip(s["big"]))
        size = 74 if n > 16 else (104 if n > 9 else 150)
        line(slide, s["big"], P(size), CLAY_LT if dark else ACC.base, bold=True,
             x=PAD_L, y=y, w=1100, h=size * 1.1, spacing=0.94)
        y += size * 0.98 + 22
    if s.get("claim"):
        line(slide, s["claim"], P(31), WHITE if dark else INK, bold=False,
             x=PAD_L, y=y, w=720, h=96, spacing=1.32)
        y += 46 + (40 if len(strip(s["claim"])) > 44 else 0)
    if s.get("note"):
        line(slide, s["note"], P(18), RGBColor(0xC8, 0xC5, 0xC0) if dark else INK_SOFT,
             x=PAD_L, y=y + 14, w=900, h=70, spacing=1.5)

def _list(slide, s, y, numbered=False):
    items = s.get("items", [])
    tight = len(items) > 5
    size = 20 if (tight or numbered) else 23
    gap = 13 if tight else 17
    avail = 596 - y
    step = min((avail / max(len(items), 1)), size * 1.42 + gap + 14)
    for i, it in enumerate(items):
        iy = y + i * step
        if numbered:
            from pptx.enum.shapes import MSO_SHAPE
            box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(PAD_L), I(iy + 1), I(28), I(28))
            box.fill.background(); box.line.color.rgb = ACC.base; box.line.width = Pt(1.1)
            box.shadow.inherit = False
            tfn = box.text_frame; tfn.word_wrap = False
            tfn.margin_left = tfn.margin_right = tfn.margin_top = tfn.margin_bottom = 0
            pn = tfn.paragraphs[0]; pn.alignment = PP_ALIGN.CENTER
            put_runs(pn, str(i + 1), P(14), ACC.base, True)
            tfn.vertical_anchor = MSO_ANCHOR.MIDDLE
            line(slide, it, P(size), INK_SOFT, x=PAD_L + 46, y=iy, w=CONTENT_W - 46,
                 h=step, spacing=1.44)
        else:
            rule(slide, PAD_L, iy + size * 0.55, 11, ACC.base, 11)
            line(slide, it, P(size), INK_SOFT, x=PAD_L + 30, y=iy, w=CONTENT_W - 30,
                 h=step, spacing=1.42)

def L_bullets(slide, s):
    bg(slide, PAPER); y = header(slide, s); _list(slide, s, y)

def L_steps(slide, s):
    bg(slide, PAPER); y = header(slide, s); _list(slide, s, y, numbered=True)

def L_twoCol(slide, s):
    bg(slide, PAPER)
    y = header(slide, s)
    colw = (CONTENT_W - 26) / 2
    from pptx.enum.shapes import MSO_SHAPE
    for xi, tone in ((PAD_L, s.get("leftTone")), (PAD_L + colw + 26, s.get("rightTone"))):
        card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(xi), I(y - 16), I(colw), I(600 - y))
        card.fill.solid()
        card.fill.fore_color.rgb = (RGBColor(0xED,0xF5,0xEE) if tone == "pos"
                                    else RGBColor(0xFB,0xED,0xEA) if tone == "neg" else ACC.t05)
        card.line.fill.background(); card.shadow.inherit = False
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(xi), I(y - 16), I(colw), I(4))
        bar.fill.solid()
        bar.fill.fore_color.rgb = POS if tone == "pos" else NEG if tone == "neg" else ACC.base
        bar.line.fill.background(); bar.shadow.inherit = False
    for side, x, head, tone in (("left", PAD_L + 26, s.get("leftHead", ""), s.get("leftTone")),
                                ("right", PAD_L + colw + 52, s.get("rightHead", ""), s.get("rightTone"))):
        c = POS if tone == "pos" else NEG if tone == "neg" else ACC.dk
        line(slide, head.upper(), P(14), c, bold=True, x=x, y=y + 10, w=colw - 52, h=24)
        items = s.get(side, [])
        step = min((580 - (y + 40)) / max(len(items), 1), 66)
        for i, it in enumerate(items):
            iy = y + 52 + i * step
            dot = POS if tone == "pos" else NEG if tone == "neg" else ACC.base
            rule(slide, x, iy + 10, 9, dot, 9)
            line(slide, it, P(18.5), INK_SOFT, x=x + 20, y=iy, w=colw - 78, h=step, spacing=1.42)

def L_table(slide, s):
    bg(slide, PAPER)
    y = header(slide, s)
    rows, head = s.get("rows", []), s.get("head")
    ncol = len(head) if head else len(rows[0])
    nrow = len(rows) + (1 if head else 0)
    size = 13.5 if len(rows) > 8 else (15.5 if len(rows) > 5 else 18)
    avail_h = 590 - y
    shape = slide.shapes.add_table(nrow, ncol, I(PAD_L), I(y), I(CONTENT_W), I(min(avail_h, nrow * 44)))
    tbl = shape.table
    tbl.first_row = bool(head)
    # column widths: first column wider
    base = CONTENT_W / ncol
    widths = [base * 1.28] + [(CONTENT_W - base * 1.28) / (ncol - 1)] * (ncol - 1) if ncol > 1 else [CONTENT_W]
    for i, w in enumerate(widths):
        tbl.columns[i].width = I(w)
    r0 = 0
    if head:
        for j, h in enumerate(head):
            cell = tbl.cell(0, j)
            cell.fill.solid(); cell.fill.fore_color.rgb = ACC.base
            cell.margin_left = cell.margin_right = I(6); cell.margin_top = cell.margin_bottom = I(5)
            tf = cell.text_frame; tf.word_wrap = True
            put_runs(tf.paragraphs[0], strip(h).upper(), P(12), WHITE, True)
        tbl.rows[0].height = I(34)
        r0 = 1
    for i, row in enumerate(rows):
        tbl.rows[i + r0].height = I(min(40, avail_h / nrow))
        for j, c in enumerate(row):
            obj = c if isinstance(c, dict) else {"t": c}
            cell = tbl.cell(i + r0, j)
            if i % 2 == 0:
                cell.fill.solid(); cell.fill.fore_color.rgb = ACC.t05
            else:
                cell.fill.background()
            cell.margin_left = cell.margin_right = I(6); cell.margin_top = cell.margin_bottom = I(5)
            cell.vertical_anchor = MSO_ANCHOR.TOP
            tf = cell.text_frame; tf.word_wrap = True
            para = tf.paragraphs[0]
            if obj.get("num"): para.alignment = PP_ALIGN.RIGHT
            col = NEG if obj.get("tone") == "neg" else POS if obj.get("tone") == "pos" else (ACC.dk if j == 0 else INK_SOFT)
            put_runs(para, obj.get("t", ""), P(size), col, bold=(j == 0 or bool(obj.get("tone"))))
    if s.get("note"):
        line(slide, s["note"], P(14.5), MUTED, x=PAD_L, y=612, w=CONTENT_W, h=40, spacing=1.45)

def L_quote(slide, s):
    dark = s.get("theme") == "dark"
    bg(slide, INK if dark else ACC.t05)
    y = 210
    if s.get("eyebrow"):
        line(slide, s["eyebrow"].upper(), P(13), MUTED, bold=True, x=PAD_L, y=150, w=900, h=24)
    q = s["quote"]
    size = 26 if len(strip(q)) > 150 else 33
    from pptx.enum.shapes import MSO_SHAPE
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(PAD_L), I(y), Emu(int(3 * 9525)), I(150))
    bar.fill.solid(); bar.fill.fore_color.rgb = ACC.base; bar.line.fill.background()
    bar.shadow.inherit = False
    line(slide, f"“{strip(q)}”", P(size), WHITE if dark else INK,
         x=PAD_L + 34, y=y - 6, w=880, h=180, spacing=1.38)
    if s.get("attrib"):
        line(slide, s["attrib"], P(16), MUTED, x=PAD_L + 37, y=y + 176, w=880, h=60, spacing=1.55)

def L_callout(slide, s):
    bg(slide, PAPER)
    y = header(slide, s)
    from pptx.enum.shapes import MSO_SHAPE
    h = 150 if s.get("detail") else 100
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(PAD_L), I(y), I(CONTENT_W), I(h))
    box.fill.solid(); box.fill.fore_color.rgb = ACC.t10
    box.line.fill.background(); box.shadow.inherit = False
    barc = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(PAD_L), I(y), I(7), I(h))
    barc.fill.solid(); barc.fill.fore_color.rgb = ACC.base
    barc.line.fill.background(); barc.shadow.inherit = False
    line(slide, s["lead"], P(27), ACC.dk, bold=True, x=PAD_L + 38, y=y + 30, w=CONTENT_W - 76,
         h=76, spacing=1.34)
    if s.get("detail"):
        line(slide, s["detail"], P(19), INK_SOFT, x=PAD_L + 38, y=y + 94, w=CONTENT_W - 76,
             h=76, spacing=1.5)

def L_kpis(slide, s):
    bg(slide, PAPER)
    y = header(slide, s)
    items = s.get("items", [])
    colw = CONTENT_W / max(len(items), 1)
    from pptx.enum.shapes import MSO_SHAPE
    for i, k in enumerate(items):
        x = PAD_L + i * colw
        card = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(x), I(y - 14), I(colw - 22), I(172))
        card.fill.solid(); card.fill.fore_color.rgb = ACC.t10
        card.line.fill.background(); card.shadow.inherit = False
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(x), I(y + 153), I(colw - 22), I(5))
        bar.fill.solid(); bar.fill.fore_color.rgb = ACC.base
        bar.line.fill.background(); bar.shadow.inherit = False
        v = str(k.get("value", ""))
        line(slide, v, P(40 if len(v) > 6 else 55), ACC.dk, bold=True,
             x=x + 22, y=y + 16, w=colw - 66, h=70)
        line(slide, k.get("label", ""), P(15), INK_SOFT,
             x=x + 22, y=y + 88, w=colw - 66, h=80, spacing=1.4)
    if s.get("note"):
        line(slide, s["note"], P(14.5), MUTED, x=PAD_L, y=y + 190, w=CONTENT_W, h=44, spacing=1.45)

def L_image(slide, s, n):
    """chart / diagram — full-bleed render from Chrome."""
    png = os.path.join(SHOTS, f"slide-{n:03d}.png")
    if os.path.exists(png):
        slide.shapes.add_picture(png, 0, 0, prs.slide_width, prs.slide_height)
    else:
        bg(slide, PAPER)
        header(slide, s)
        line(slide, "[figure missing]", P(20), MUTED, x=PAD_L, y=340, w=600, h=40)

def L_explain(slide, s):
    """The idea in plain words: story on the left, use cases + buyers on the right."""
    bg(slide, PAPER)
    y = header(slide, s)
    split = PAD_L + CONTENT_W * 0.59
    left_w, xr = split - PAD_L - 44, split + 44
    right_w = PAD_L + CONTENT_W - xr
    vrule(slide, split, y, 600 - y)
    line(slide, s.get("mainHead", "What it does").upper(), P(14), ACC.dk, bold=True,
         x=PAD_L, y=y, w=left_w, h=24)
    tl = line(slide, s.get("story", ""), P(21), INK, x=PAD_L, y=y + 38, w=left_w, h=600 - y - 38, spacing=1.45)
    if s.get("why"):
        line(slide, s.get("whyHead", "Why it matters").upper(), P(14), ACC.dk, bold=True, tf=tl,
             space_before=18, space_after=8)
        line(slide, s["why"], P(21), INK, tf=tl, spacing=1.45)
    tf = textbox(slide, xr, y, right_w, 600 - y)
    line(slide, s.get("listHead", "Use cases").upper(), P(14), ACC.dk, bold=True, tf=tf, space_after=10)
    for it in s.get("uses", []):
        line(slide, "–  " + it, P(18), INK_SOFT, tf=tf, space_after=8, spacing=1.4)
    line(slide, s.get("buyersHead", "Who buys it").upper(), P(14), ACC.dk, bold=True, tf=tf,
         space_before=16, space_after=8)
    line(slide, s.get("buyers", ""), P(18), INK_SOFT, tf=tf, spacing=1.4)

LAYOUTS = {
    "explain": L_explain,
    "title": L_title, "section": L_section, "statement": L_statement,
    "bullets": L_bullets, "steps": L_steps, "twoCol": L_twoCol, "table": L_table,
    "quote": L_quote, "callout": L_callout, "kpis": L_kpis,
}

# ---------------- build ----------------

files = sorted(f for f in os.listdir(SLIDES_DIR) if f.endswith(".json"))
all_slides = []
for f in files:
    data = json.loads(Path(SLIDES_DIR, f).read_text())
    arr = data if isinstance(data, list) else data.get("slides", [])
    for _s in arr:
        _s["_file"] = f
    all_slides.extend(arr)

made = 0
for i, s in enumerate(all_slides, 1):
    slide = prs.slides.add_slide(BLANK)
    lay = s.get("layout")
    ACC.hex = ACCENTS[(int(str(s.get("_file", "00"))[:2] or 0)) % len(ACCENTS)]
    if lay in ("chart", "diagram"):
        L_image(slide, s, i)
    elif lay in LAYOUTS:
        LAYOUTS[lay](slide, s)
        if lay not in ("title",):
            furniture(slide, s, i, dark=(s.get("theme") == "dark" or lay == "section"))
    else:
        bg(slide, PAPER)
        line(slide, f"[unknown layout: {lay}]", P(20), NEG, x=PAD_L, y=340, w=800, h=40)
    if lay not in ("title", "section", "chart", "diagram"):
        from pptx.enum.shapes import MSO_SHAPE
        topbar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, I(7))
        topbar.fill.solid(); topbar.fill.fore_color.rgb = ACC.base
        topbar.line.fill.background(); topbar.shadow.inherit = False
    made += 1

prs.save(OUT)
print(f"{made} slides -> {OUT}")
